from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from typing import Literal

import pdfplumber
from docx import Document

MAX_RESUME_BYTES = 5 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_DOCX_ENTRIES = 1000
MIN_EXTRACTED_CHARACTERS = 200

ResumeFileType = Literal["pdf", "docx", "txt"]


class ResumeParseError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class ParsedResume:
    text: str
    file_type: ResumeFileType
    character_count: int


def parse_resume(data: bytes) -> ParsedResume:
    if not data:
        raise ResumeParseError("INVALID_FILE", "The uploaded file is empty.")
    if len(data) > MAX_RESUME_BYTES:
        raise ResumeParseError(
            "FILE_TOO_LARGE",
            "Resume files must be 5 MB or smaller.",
        )

    pdf_header = data.find(b"%PDF-")
    if 0 <= pdf_header <= 1024:
        file_type: ResumeFileType = "pdf"
        text = _extract_pdf(data)
    elif data.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")):
        file_type = "docx"
        text = _extract_docx(data)
    elif data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        raise ResumeParseError(
            "UNSUPPORTED_FILE_TYPE",
            "This legacy Office format is not supported. Upload PDF, DOCX, or TXT.",
        )
    else:
        file_type = "txt"
        text = _extract_txt(data)

    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if len(text) < MIN_EXTRACTED_CHARACTERS:
        if file_type == "pdf":
            raise ResumeParseError(
                "SCANNED_PDF",
                "This PDF has too little extractable text. Upload a text-based PDF "
                "or DOCX; OCR is not supported.",
            )
        raise ResumeParseError(
            "INSUFFICIENT_TEXT",
            "This file has too little readable text. Upload a text-based resume "
            "with at least 200 characters.",
        )

    return ParsedResume(text=text, file_type=file_type, character_count=len(text))


def _extract_pdf(data: bytes) -> str:
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if getattr(pdf.doc, "encryption", None) is not None:
                raise ResumeParseError(
                    "ENCRYPTED_PDF",
                    "Encrypted PDFs are not supported. Upload an unlocked PDF or DOCX.",
                )
            pages = [page.extract_text() or "" for page in pdf.pages]
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError(
            "INVALID_FILE",
            "The PDF could not be read. Upload a valid, unlocked PDF.",
        ) from exc
    return "\n".join(pages)


def _extract_docx(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_DOCX_ENTRIES:
                raise ResumeParseError(
                    "INVALID_FILE",
                    "The DOCX archive contains too many files.",
                )
            if sum(entry.file_size for entry in entries) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ResumeParseError(
                    "FILE_TOO_LARGE",
                    "The DOCX expands beyond the 20 MB safety limit.",
                )
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise ResumeParseError(
                    "ENCRYPTED_FILE",
                    "Encrypted DOCX files are not supported.",
                )

            names = {entry.filename.lower() for entry in entries}
            if "word/document.xml" not in names:
                raise ResumeParseError(
                    "INVALID_FILE",
                    "The uploaded ZIP file is not a valid DOCX document.",
                )
            if any("vbaproject.bin" in name for name in names):
                raise ResumeParseError(
                    "MACROS_NOT_SUPPORTED",
                    "Macro-enabled Word documents are not supported.",
                )
            content_types = archive.read("[Content_Types].xml").lower()
            if b"macroenabled" in content_types:
                raise ResumeParseError(
                    "MACROS_NOT_SUPPORTED",
                    "Macro-enabled Word documents are not supported.",
                )

        document = Document(io.BytesIO(data))
    except ResumeParseError:
        raise
    except (KeyError, OSError, zipfile.BadZipFile, ValueError) as exc:
        raise ResumeParseError(
            "INVALID_FILE",
            "The DOCX could not be read. Upload a valid DOCX document.",
        ) from exc

    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_rows = [
        "\t".join(cell.text for cell in row.cells)
        for table in document.tables
        for row in table.rows
    ]
    return "\n".join((*paragraphs, *table_rows))


def _extract_txt(data: bytes) -> str:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ResumeParseError(
            "UNSUPPORTED_FILE_TYPE",
            "The file is not a supported PDF, DOCX, or UTF-8 text document.",
        ) from exc
    if "\x00" in text:
        raise ResumeParseError(
            "UNSUPPORTED_FILE_TYPE",
            "The file is not a supported PDF, DOCX, or UTF-8 text document.",
        )
    return text
